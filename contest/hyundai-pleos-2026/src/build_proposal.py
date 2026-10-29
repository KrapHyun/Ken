# -*- coding: utf-8 -*-
"""공모전 접수양식(02 앱 프로젝트 소개)에 기획서 초안을 채워 넣는다.

사용법: python3 src/build_proposal.py
입력: src/application-form-blank.docx, figures/*.png
출력: PhonePilot_기획서_초안.docx
"""
import copy
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FORM = ROOT / "src" / "application-form-blank.docx"
FIG = ROOT / "figures"
OUT = ROOT / "PhonePilot_기획서_초안.docx"

APP_NAME = "PhonePilot (폰파일럿) – 거치만 하면 내 폰이 차의 눈과 두뇌가 된다"
NAVY = RGBColor(0x00, 0x2C, 0x5F)
GREY = RGBColor(0x5B, 0x64, 0x74)
FIG_WIDTH = Cm(16.0)


# ---------------------------------------------------------------- helpers
def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


class Box:
    """양식의 1x1 입력 박스(셀)에 내용을 순서대로 추가한다."""

    def __init__(self, table):
        self.cell = table.cell(0, 0)
        self._first = self.cell.paragraphs[0]
        self._used_first = False

    def _para(self):
        if not self._used_first:
            self._used_first = True
            return self._first
        return self.cell.add_paragraph()

    def p(self, text="", bold=False, size=10, color=None, align=None, space_after=3):
        para = self._para()
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(0)
        if align:
            para.alignment = align
        # "**굵게**" 구간을 지원
        parts = text.split("**")
        for i, part in enumerate(parts):
            if not part:
                continue
            r = para.add_run(part)
            r.bold = bold or (i % 2 == 1)
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return para

    def h(self, text):
        para = self.p(text, bold=True, size=11, color=NAVY, space_after=2)
        para.paragraph_format.space_before = Pt(6)
        return para

    def bullet(self, text, level=0):
        para = self.p(("    " * level) + "• " + text, size=10, space_after=1)
        para.paragraph_format.left_indent = Cm(0.4 + 0.5 * level)
        para.paragraph_format.first_line_indent = Cm(-0.35)
        return para

    def note(self, text):
        return self.p(text, size=8.5, color=GREY, space_after=3)

    def figure(self, name, caption):
        para = self._para()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(4)
        para.add_run().add_picture(str(FIG / name), width=FIG_WIDTH)
        self.p(caption, size=8.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

    def table(self, header, rows, widths, font=8.5):
        # 셀 안의 표는 앞 단락 뒤에 붙는다
        self._para()
        t = self.cell.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(Cm(w).twips)))
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Cm(w)
        for ci, text in enumerate(header):
            c = t.rows[0].cells[ci]
            _shade(c, "002C5F")
            c.paragraphs[0].text = ""
            r = c.paragraphs[0].add_run(text)
            r.bold = True
            r.font.size = Pt(font)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for ri, row in enumerate(rows, start=1):
            for ci, text in enumerate(row):
                c = t.rows[ri].cells[ci]
                if ri % 2 == 0:
                    _shade(c, "EEF3FA")
                para = c.paragraphs[0]
                parts = str(text).split("**")
                for i, part in enumerate(parts):
                    if part:
                        r = para.add_run(part)
                        r.bold = i % 2 == 1
                        r.font.size = Pt(font)
        # 표 뒤 여백용 빈 단락
        tail = self.cell.add_paragraph()
        tail.paragraph_format.space_after = Pt(2)
        return t


def fill_blank_line(paragraph, value):
    """'앱명 / 프로젝트명: ______' 줄의 밑줄을 값으로 바꾼다."""
    text = paragraph.text
    head = text.split(":")[0]
    for r in paragraph.runs[1:]:
        r.text = ""
    paragraph.runs[0].text = f"{head}: {value}"


# ---------------------------------------------------------------- content
def build():
    d = docx.Document(str(FORM))
    for para in d.paragraphs:
        if para.text.strip().startswith("앱명 / 프로젝트명"):
            fill_blank_line(para, "PhonePilot (폰파일럿)")

    # 긴 박스가 페이지를 넘어 나뉘도록 행 분할 금지 해제
    for t in d.tables[1:9]:
        for el in list(t._tbl.iter(qn("w:cantSplit"))):
            el.getparent().remove(el)
    boxes = [Box(d.tables[i]) for i in range(1, 9)]
    b1, b2, b3, b4, b5, b6, b7, b8 = boxes

    # 1. 한 줄 소개 ---------------------------------------------------------
    b1.p("차량 거치대에 올린 스마트폰을 Pleos Connect 차량의 **추가 센서·엣지 AI**로 연결해, "
         "운전자 주의 모니터링으로 **주행 안전**을 높이고, 수집된 주행 데이터의 품질만큼 "
         "**리워드(포인트·구독 할인)**를 돌려주는 차량-스마트폰 연동 앱", size=10.5)

    # 2. 기획 배경 ----------------------------------------------------------
    b2.h("① 고객의 불편한 순간: \"크루즈 켜두면 자꾸 폰을 보게 된다\"")
    b2.p("운전 보조(ACC·LKA) 기능이 편해질수록 운전자는 도로에서 눈을 떼기 쉽습니다. "
         "미국 IIHS·MIT 공동 연구에서 ACC+차로유지 보조를 4주간 사용한 운전자는 사용 기간이 길어질수록 "
         "휴대폰 조작 등 비주행 행동이 늘어났습니다 [3]. 미국에서는 2023년 한 해 주의분산 운전으로 "
         "3,275명이 사망했고 [1], 사고의 결정적 원인 중 약 94%가 운전자 요인으로 분석됩니다 [2].")
    b2.h("② 규제와 시장의 흐름: 운전자 모니터링(DMS)은 '필수'가 되는 중")
    b2.p("EU 일반안전규정(GSR, 2019/2144)은 졸음 경고(DDAW)와 고도 주의분산 경고(ADDW)를 신차에 "
         "단계적으로 의무화했습니다 [4]. 그러나 DMS 카메라는 상위 트림·옵션 위주로 탑재되어, "
         "많은 운전자가 여전히 이 보호를 받지 못합니다.")
    b2.h("③ 제조사의 고민: 자율주행 학습 데이터는 비싸다")
    b2.p("자율주행·ADAS 성능은 다양한 실도로 데이터의 양과 질에 좌우됩니다. 전용 수집 차량을 운영하는 "
         "비용은 크고, 엣지 케이스(급제동, 끼어들기, 공사구간)는 드물게 발생합니다. 스마트폰급 기기로 "
         "영상·CAN·GNSS·IMU를 함께 모은 comma2k19 [6], 대시캠 앱으로 10만 개 주행 영상을 크라우드소싱한 "
         "BDD100K [7]는 '스마트폰 기반 수집'이 학습 데이터로서 가치가 있음을 보여줍니다.")
    b2.h("④ 인사이트: 운전자는 이미 매일 '고성능 센서 묶음'을 들고 탄다")
    b2.p("최신 스마트폰에는 다중 카메라, 9축 IMU, GPS, 기압계, NPU/GPU가 모두 들어 있고, 많은 운전자가 "
         "이미 거치대에 폰을 올려 둡니다. PhonePilot은 이 폰을 **차량과 연동되는 센서·연산 장치**로 바꿔, "
         "하나의 행동(거치)으로 세 가지 문제를 동시에 해결합니다.")
    b2.table(["문제", "현재", "PhonePilot의 해결"], [
        ["운전 중 폰 사용·주의분산", "운전자 의지에 의존", "폰은 고정된 채 연산 중(화면은 시계만) + 전면 카메라로 전방주시 감시·경고"],
        ["DMS 미탑재 차량", "옵션·상위 트림 위주", "폰 전면 카메라 기반 졸음·주의분산 경고를 앱으로 제공"],
        ["자율주행 데이터 확보 비용", "전용 수집 차량·인력", "고객 차량 + 고객 폰 = 분산 수집 플릿, 품질 기반 리워드로 참여 유도"],
        ["구독 서비스 진입 장벽", "유료 구독 망설임", "데이터 기여만큼 구독 할인·기능 언락·포인트 → 자연스러운 구독 경험"],
    ], widths=[3.6, 3.6, 9.2])
    b2.note("※ 인용 수치는 각 기관 공개 자료 기준이며, 제출 전 최신 발표치로 재확인 예정.")

    # 3. 주요 기능 ----------------------------------------------------------
    b3.p("PhonePilot은 **스마트폰 컴패니언 앱**(센서·온디바이스 AI)과 **Pleos Connect 차량 앱**"
         "(차량 신호·공조 제어·UI)이 한 쌍으로 동작합니다.")
    b3.figure("fig1.png", "그림 1. 시스템 구조")
    b3.table(["#", "핵심 기능", "해결하는 문제 / 새로운 경험", "차량 특화 요소 (Pleos SDK)"], [
        ["F1", "**AI 거치 검사 (Mount Check)**\n후면=도로, 전면=얼굴 화각 자동 판정 후 OK 시 자동 연동",
         "누구나 '좋은 위치'에 거치 → 데이터 품질 확보, 설치 가이드 제공", "연동 성공 시 차량 화면에 상태 표시"],
        ["F2", "**동기화 주행 데이터 수집**\n전방 영상 + IMU 100Hz + GPS + 기압 + 차량 신호",
         "영상에 '운전자가 실제로 한 조작'이 정답 라벨로 붙은 학습 데이터", "Vehicle SDK: 차속·조향각·브레이크·기어·크루즈 상태"],
        ["F3", "**운전자 주의 모니터링 (DMS)**\n눈 감김(EAR)·PERCLOS·2초 이상 시선 이탈 [5] 감지",
         "크루즈/LKA 사용 중 주의분산 사고 예방", "Gleo AI 음성 경고, ADAS 작동 중 경고 민감도 상향"],
        ["F4", "**이상 이벤트 자동 클립**\n급제동·급조향·ADAS 해제 등 전후 15초 저장",
         "가치 있는 장면만 선별 → 전송량 최소화", "차량 신호 기반 이벤트 트리거(차속·조향 변화)"],
        ["F5", "**폰 발열·전원 관리**\n발열 단계별 공조 보조 → 프레임 감소 → 안전 종료",
         "장시간 수집 시 과열·성능 저하 방지", "HVAC 제어: 거치 송풍구 냉방 보조, 히터 토출 회피"],
        ["F6", "**주행 중 방해 차단**\n수집 중 폰 화면 OFF(시계만), 전화는 차량 핸즈프리",
         "폰을 '볼 수 없는' 환경을 자연스럽게 형성", "차량 디스플레이에 요약 정보만 표시(DDO 준수)"],
        ["F7", "**비식별화 + 품질 기반 리워드**\n얼굴·번호판 블러, AI 품질 점수 × 유효 km",
         "프라이버시 보호 + 공정한 보상, 부정 수집 차단", "영상-차속 정합 검사로 부정 수집 판별"],
    ], widths=[0.8, 5.0, 5.4, 5.2], font=8)
    b3.figure("fig2.png", "그림 2. 핵심 알고리즘 파이프라인")
    b3.h("센서별 수집 데이터와 활용처")
    b3.table(["센서", "수집 데이터", "학습·서비스 활용"], [
        ["후면 카메라", "전방 도로 영상 (1080p 30fps, OIS 고정)", "주행 환경 인식, 경로 분석, 자율주행 인지·주행 정책 학습 [8][12]"],
        ["전면 카메라", "운전자 얼굴 랜드마크 지표 (원본 미저장)", "전방주시·졸음 지표 [9][10][11], 운전 피로도·만족도 분석"],
        ["IMU (가속도·자이로·지자기)", "3축×3 100Hz", "차량 거동(종·횡가속, 요레이트) 학습, 노면 상태 추정"],
        ["GPS · 기압계", "위치·속도·고도", "위치 맥락, 경사로·고가도로 판별, 측위 보정"],
        ["차량 신호 (Pleos)", "차속, 조향각, 브레이크, 기어, 크루즈/ADAS 상태", "영상과 정렬된 정답 라벨(행동 복제·섀도 모드 비교 기준)"],
    ], widths=[3.2, 5.6, 7.6])

    # 4. 사용 시나리오 ------------------------------------------------------
    b4.h("페르소나 A — 고속도로 출퇴근 직장인 (35세, 하루 왕복 70km, 크루즈 상시 사용)")
    b4.table(["시점", "사용자 행동", "스마트폰", "Pleos 차량 앱"], [
        ["탑승", "폰을 송풍구 거치대에 올리고 USB 연결", "PhonePilot 자동 실행 → 후면·전면 카메라 화각 검사",
         "\"폰이 연결되었습니다\" 표시"],
        ["거치 검사", "화면 안내에 따라 각도 조정", "기울기·지평선·노면 비율·얼굴 인식 → 점수 92, OK",
         "수집 시작 + 해당 송풍구 냉방 보조, 히터 토출 회피"],
        ["주행", "크루즈 켜고 주행", "화면 OFF(시계), 영상·센서 동기 기록, DMS 동작",
         "수집 중 배지, 오늘의 유효 km"],
        ["위험 순간", "졸음으로 눈 감김 2초", "PERCLOS 상승 감지 → 경고 요청", "Gleo AI 음성: \"휴식이 필요해요. 가까운 졸음쉼터를 안내할까요?\""],
        ["이벤트", "앞차 급제동 → 급감속", "전후 15초 클립 저장(얼굴·번호판 블러)", "이벤트 1건 기록"],
        ["하차 후", "귀가, 폰 충전", "새벽 + Wi-Fi + 충전 중 자동 업로드", "—"],
        ["다음 날", "리워드 확인", "품질 점수 0.87 → 포인트 적립", "구독 할인율·포인트 표시"],
    ], widths=[1.8, 3.8, 5.5, 5.3], font=8)
    b4.h("페르소나 B — 영업직 장거리 운전자 (월 3,000km)")
    b4.p("주행거리가 많아 리워드 체감이 큽니다. 적립 포인트로 커넥티드 서비스 구독을 무료로 유지하고, "
         "피로 누적 구간에서 DMS 경고와 휴식 안내를 받습니다.")
    b4.h("페르소나 C — 얼리어답터 (섀도 모드 베타 참가자)")
    b4.p("전용 거치대를 장착해 고정 캘리브레이션 데이터를 제공하고, 가산 리워드와 신기능 선공개 혜택을 받습니다. "
         "주행 후 \"오늘 AI가 내 운전과 몇 % 같았나\" 리포트를 확인합니다.")
    b4.figure("fig4.png", "그림 4. 화면 시안 – (좌) 스마트폰 거치 검사, (우) Pleos Connect 차량 디스플레이")

    # 5. 개발 계획 ----------------------------------------------------------
    b5.h("① Pleos Connect SDK/API 사용 계획")
    b5.table(["SDK", "사용 API·데이터", "목적", "MVP"], [
        ["Vehicle SDK", "차속, 조향각, 브레이크, 기어 (권한 선언 후 구독)", "영상 정답 라벨, 이벤트 트리거, 부정 수집 판별", "●"],
        ["Vehicle SDK – DrivingMode", "크루즈 컨트롤 작동 상태", "보조 주행 중 DMS 민감도 상향", "●"],
        ["Vehicle SDK – HVAC", "공조 상태 조회/설정", "폰 발열 시 냉방 보조, 거치 송풍구 히터 회피", "●"],
        ["ADAS SDK", "ADAS 작동·해제 상태", "해제 순간 이벤트 클립, 섀도 모드 비교", "○"],
        ["Vehicle Location SDK", "차량 위치", "폰 GPS 보정·위치 맥락", "○"],
        ["Gleo AI SDK", "음성 안내·명령", "졸음 경고, \"수집 시작/중지\" 음성 명령", "○"],
        ["Pleos 에뮬레이터 (VHAL)", "차량 신호 시뮬레이션", "실차 없이 연동 개발·시연", "●"],
    ], widths=[3.6, 4.6, 6.4, 1.8], font=8)
    b5.note("● 1주 MVP 구현 / ○ 본선(12월) 이후 확장. 세부 API 명세·권한 범위는 Pleos Playground 공식 문서 기준으로 확인 후 적용 [15].")
    b5.h("② 폰-차량 연동 방식")
    b5.bullet("1안: 로컬 연결(USB/Wi-Fi)로 차량 신호를 폰에 실시간 전달")
    b5.bullet("2안(폴백): 차량 앱과 폰 앱이 각각 단조 시계 타임스탬프로 기록 → 업로드 후 서버에서 병합. "
              "폰 IMU 종가속도와 차속 미분값의 상호상관으로 시간 오프셋 자동 보정(수십 ms 수준 목표)")
    b5.h("③ 기술 스택")
    b5.table(["영역", "기술"], [
        ["Pleos 차량 앱", "Kotlin, Android Automotive OS, Pleos Connect SDK, Driver Distraction Optimized UI"],
        ["스마트폰 앱", "Kotlin, CameraX/Camera2(동시 카메라 API, 미지원 기기는 전·후면 교대 촬영), SensorManager, "
                     "MediaPipe Face Landmarker [9], OpenCV, WorkManager(새벽·Wi-Fi·충전 조건), PowerManager 발열 상태 API"],
        ["온디바이스 AI", "얼굴 랜드마크 → EAR [10]·PERCLOS [11]·머리 자세, 경량 객체 검출(Apache/MIT 라이선스 모델 우선), 얼굴·번호판 블러"],
        ["분석 도구", "Python(pandas, OpenCV) – 동기화 검증, 이벤트 시각화, 품질 점수 산정 노트북"],
    ], widths=[3.2, 13.2])
    b5.h("④ 1주 MVP 개발 일정 (공개 GitHub 저장소로 제출)")
    b5.table(["일차", "작업", "산출물"], [
        ["Day 1", "폰 앱 골격, 후면 카메라 녹화 + IMU/GPS/기압 CSV 로깅(단조 시계)", "raw 로그 1회분"],
        ["Day 2", "Mount Check: 중력벡터 기반 롤·피치, 예상 지평선 오버레이, 선명도·노출 검사, OK 판정", "거치 검사 화면"],
        ["Day 3", "전면 카메라 얼굴 랜드마크 → EAR·시선 이탈 경고, 화면 OFF(시계) 모드", "DMS 데모"],
        ["Day 4", "IMU 규칙 기반 이벤트 감지 + 링버퍼 클립 저장, 발열 상태 모니터링", "이벤트 클립"],
        ["Day 5", "Pleos 차량 앱(에뮬레이터): 차속·조향·크루즈 구독, 수집 상태 UI, HVAC 보조 호출", "차량 앱 APK"],
        ["Day 6", "실차 주행 수집 → 비식별화 → 샘플 데이터셋 정리, Python 동기화·시각화 노트북", "sample-data/, notebook"],
        ["Day 7", "시연 영상(거치 → 검사 → OK → 연동 시작), README·데이터 명세 작성", "영상, GitHub 공개"],
    ], widths=[1.6, 10.6, 4.2], font=8)
    b5.h("⑤ 샘플 데이터 명세 (GitHub 공개분)")
    b5.table(["파일", "내용", "주기"], [
        ["front_video.mp4", "전방 도로 영상(얼굴·번호판 블러 처리)", "30fps"],
        ["imu.csv", "t_ns, ax, ay, az, gx, gy, gz, mx, my, mz", "100Hz"],
        ["gnss.csv / baro.csv", "t_ns, lat, lon, speed, bearing, acc / t_ns, hPa", "1–10Hz / 25Hz"],
        ["vehicle.csv", "t_ns, speed_kph, steer_deg, brake, gear, cruise (에뮬레이터·실차)", "10Hz"],
        ["dms.csv", "t_ns, ear_l, ear_r, perclos_60s, head_yaw, head_pitch, alert", "15Hz"],
        ["events.json / meta.json", "이벤트 유형·시각·클립 경로 / 기기·거치 검사 점수·캘리브레이션", "이벤트 시"],
    ], widths=[4.0, 9.6, 2.8], font=8)
    b5.h("⑥ 개인정보·안전 설계")
    b5.table(["항목", "설계"], [
        ["운전자 얼굴", "원본 영상 미저장, 온디바이스에서 지표(EAR·시선)만 산출. 감정·만족도 분석은 별도 선택 동의"],
        ["차 밖 보행자·번호판", "업로드 전 온디바이스 자동 블러 (개인정보보호법상 이동형 영상정보처리기기 규정 대응)"],
        ["위치정보", "위치정보 수집 동의 + 운영 주체(제조사)의 위치기반서비스 사업 요건 하에 처리"],
        ["주행 안전", "수집 중 폰 화면 OFF, 차량 화면은 요약 정보만, 경고는 음성·진동 우선"],
        ["리워드 부정 방지", "영상 움직임-차속-IMU 정합 검사, 거치 검사 점수 하한, 품질 점수 기반 지급"],
    ], widths=[3.4, 13.0])

    # 6. 기대효과 -----------------------------------------------------------
    b6.h("① 사용자 가치")
    b6.bullet("**안전**: 크루즈·LKA 사용 중에도 전방주시를 지켜주는 DMS를 하드웨어 추가 없이 제공")
    b6.bullet("**보상**: 평소처럼 운전만 해도 품질 점수만큼 포인트·구독 할인 → 구독 서비스를 무료/저가로 경험")
    b6.bullet("**편의**: 거치 → 자동 검사 → 자동 연동 → 새벽 자동 업로드, 추가 조작 0회")
    b6.h("② 현대자동차그룹 가치")
    b6.bullet("**데이터 플릿**: Pleos Connect 탑재 차량(2030년 2,000만 대 이상 목표 [13])이 곧 분산 데이터 수집 플릿")
    b6.bullet("**정답 라벨 데이터**: 영상 + 차량 신호가 정렬된 학습 데이터 → 인지·주행 정책 학습 [8], 섀도 모드 검증")
    b6.bullet("**구독 전환**: 리워드로 구독 서비스 진입 장벽을 낮추고 유지율 향상")
    b6.bullet("**신규 매출**: 냉각·고정 캘리브레이션 전용 거치대 1st party 액세서리, 보험사 UBI(운전 습관 연계 보험) 제휴")
    b6.h("③ 타깃 고객 및 예상 MAU (가정 기반 추정)")
    b6.table(["구분", "1년차", "2년차", "3년차"], [
        ["Pleos Connect 탑재 운행 차량 (가정)", "30만 대", "120만 대", "300만 대"],
        ["앱 설치율 (가정)", "8%", "10%", "12%"],
        ["월간 활성 비율 (가정)", "60%", "60%", "60%"],
        ["**예상 MAU**", "**약 1.4만**", "**약 7.2만**", "**약 21.6만**"],
        ["월 수집 주행거리 (MAU × 800km)", "약 1,150만 km", "약 5,760만 km", "약 1.7억 km"],
    ], widths=[6.4, 3.3, 3.3, 3.4])
    b6.note("※ 1차 타깃: 크루즈·고속도로 주행 비중이 높은 출퇴근·장거리 운전자, 2차: 데이터 리워드에 반응하는 2030 운전자. "
            "탑재 차량 수는 공개 목표치를 바탕으로 한 보수적 가정이며 실제 판매 데이터로 보정 필요.")
    b6.h("④ Retention 전략 (목표: D30 45%, 6개월 30%)")
    b6.bullet("거치만 하면 자동 시작(무노력 습관화), 주행마다 쌓이는 포인트(가변 보상), 주간 안전 리포트")
    b6.bullet("품질 등급제(브론즈→골드): 좋은 위치·전용 거치대 사용 시 가산 → 장기 참여와 데이터 품질 동시 상승")
    b6.h("⑤ 향후 확대 계획")
    b6.figure("fig3.png", "그림 3. 단계별 로드맵 (데이터 수집 → 섀도 모드 → 베타 → 적용)")
    b6.p("장기적으로는 폰-차량 연산 분담, 그리고 폰 상태 이상감지 시 운전자에게 사전 제어권 전환 요청(TOR)을 보내고 "
         "응답이 없으면 최소위험기동(MRM, 갓길 정차)으로 이어지는 구조로 확장합니다. 이 단계는 기능안전(ISO 26262)·"
         "관련 법규 인증 범위 안에서 제조사와 함께 단계적으로 검증합니다.")

    # 7. 참고자료 ----------------------------------------------------------
    b7.h("제출 자료")
    b7.bullet("GitHub 저장소: https://github.com/<계정>/phonepilot (폰 앱·Pleos 차량 앱 소스, 샘플 데이터, 분석 노트북) — 공개 후 기재")
    b7.bullet("시연 영상: 폰 거치 → 후면 카메라 도로 화각·전면 카메라 얼굴 검사 → OK → 차량 연동 시작 (약 1분) — 업로드 후 기재")
    b7.bullet("앱 화면 시안: 본문 그림 4, 시스템 구조·알고리즘·로드맵: 그림 1~3")
    b7.h("참고문헌·자료")
    refs = [
        "[1] NHTSA, \"Distracted Driving in 2023\", Traffic Safety Facts Research Note, 2025.",
        "[2] S. Singh, \"Critical Reasons for Crashes Investigated in the National Motor Vehicle Crash Causation Survey\", NHTSA DOT HS 812 115, 2015.",
        "[3] I. J. Reagan et al., \"Disengagement from driving when using automation during a 4-week field trial\", Transportation Research Part F, 82, 2021.",
        "[4] Regulation (EU) 2019/2144 (General Safety Regulation) – DDAW·ADDW 요건.",
        "[5] NHTSA, \"Visual-Manual NHTSA Driver Distraction Guidelines for In-Vehicle Electronic Devices\", 2013 (시선 이탈 2초 기준).",
        "[6] H. Schafer, E. Santana, A. Haden, R. Biasini, \"A Commute in Data: The comma2k19 Dataset\", arXiv:1812.05752, 2018.",
        "[7] F. Yu et al., \"BDD100K: A Diverse Driving Dataset for Heterogeneous Multitask Learning\", CVPR 2020.",
        "[8] M. Bojarski et al., \"End to End Learning for Self-Driving Cars\", arXiv:1604.07316, 2016.",
        "[9] Y. Kartynnik et al., \"Real-time Facial Surface Geometry from Monocular Video on Mobile GPUs\", arXiv:1907.06724, 2019.",
        "[10] T. Soukupová, J. Čech, \"Real-Time Eye Blink Detection using Facial Landmarks\", CVWW 2016.",
        "[11] D. F. Dinges, R. Grace, \"PERCLOS: A Valid Psychophysiological Measure of Alertness\", FHWA-MCRT-98-006, 1998.",
        "[12] Z. Qin, H. Wang, X. Li, \"Ultra Fast Structure-aware Deep Lane Detection\", ECCV 2020.",
        "[13] 현대자동차그룹, Pleos 브랜드 및 Pleos Connect 발표 보도자료 (Pleos 25, 2025).",
        "[14] 현대자동차그룹, \"현대차·기아, 차량용 앱 개발 공모전 개최\", 2026.10.",
        "[15] Pleos Playground 개발자 문서 – Vehicle SDK, Safe Driving Score 튜토리얼, Pleos Connect 에뮬레이터 (document.pleos.ai).",
    ]
    for r in refs:
        b7.p(r, size=8.5, space_after=1)

    # 8. 숏폼 --------------------------------------------------------------
    b8.p("(선택) 거치 → OK → 연동 시작 과정을 담은 세로 영상 링크 기재 예정", size=9.5, color=GREY)

    d.save(str(OUT))
    print("saved", OUT)


if __name__ == "__main__":
    build()
